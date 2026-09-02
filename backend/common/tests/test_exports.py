from io import BytesIO

import pytest
from django.test import override_settings
from openpyxl import load_workbook

from audit.models import AuditLog

pytestmark = pytest.mark.django_db

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture
def admin_client(user_factory, as_user):
    return as_user(user_factory(roles=["Admin"]))


def _rows(response):
    workbook = load_workbook(BytesIO(response.content))
    return list(workbook.active.iter_rows(values_only=True))


@pytest.mark.parametrize(
    "path,action",
    [
        ("/api/v1/users/export/", "user.export"),
        ("/api/v1/roles/export/", "role.export"),
        ("/api/v1/audit-logs/export/", "audit.export"),
    ],
)
def test_export_returns_xlsx_attachment_and_audits(admin_client, path, action):
    resp = admin_client.get(path)
    assert resp.status_code == 200
    assert resp["Content-Type"] == XLSX_MIME
    assert "attachment;" in resp["Content-Disposition"]
    assert resp["Content-Disposition"].endswith('.xlsx"')

    entry = AuditLog.objects.get(action=action)
    assert entry.changes["format"] == "xlsx"
    assert "rows" in entry.changes


@override_settings(EXPORT_MAX_ROWS=2)
def test_export_respects_the_row_cap(admin_client, user_factory):
    for i in range(5):
        user_factory(email=f"cap-{i}@example.com")

    resp = admin_client.get("/api/v1/users/export/")
    data_rows = _rows(resp)[1:]  # drop header
    assert len(data_rows) == 2
    assert AuditLog.objects.get(action="user.export").changes["rows"] == 2


def test_export_applies_the_active_filters(admin_client, user_factory):
    user_factory(email="active-export@example.com", is_active=True)
    user_factory(email="inactive-export@example.com", is_active=False)

    resp = admin_client.get("/api/v1/users/export/?is_active=false")
    emails = {row[1] for row in _rows(resp)[1:]}
    assert emails == {"inactive-export@example.com"}


def test_export_permission_is_enforced(user_factory, as_user):
    viewer = as_user(user_factory(roles=["Viewer"]))
    assert viewer.get("/api/v1/audit-logs/export/").status_code == 403
    assert viewer.get("/api/v1/users/export/").status_code == 200
